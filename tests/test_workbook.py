import csv

import pandas as pd
import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from pbi_interaction.batch import FIELDS, remove_csvs
from pbi_interaction.settings import SheetSpec, WorkbookSpec
from pbi_interaction.workbook import build_workbooks, merge_frames, plan_workbook


def wb(*sheets, unassigned="sheet"):
    return WorkbookSpec(name="w", sheets=list(sheets), unassigned=unassigned)


def sheet(name, runs, **kw):
    return SheetSpec(name=name, runs=runs, **kw)


def make_batch(tmp_path, runs: dict[str, tuple[str, str]]):
    """runs: label -> (csv text, status)"""
    with (tmp_path / "manifest.csv").open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        for label, (text, st) in runs.items():
            (tmp_path / f"{label}.csv").parent.mkdir(parents=True, exist_ok=True)
            (tmp_path / f"{label}.csv").write_text(text)
            w.writerow({f: "" for f in FIELDS} | {"label": label, "status": st, "file": f"{label}.csv"})


def test_plan_globs_unassigned_and_order():
    labels = ["a/x", "a/y", "b", "c"]
    plans = plan_workbook(wb(sheet("A", ["a/*"], mode="aggregate", key=["k"]), sheet("B", ["b"])), labels)
    assert [(p.name, p.labels) for p in plans] == [("A", ["a/x", "a/y"]), ("B", ["b"]), ("c", ["c"])]


def test_plan_errors():
    with pytest.raises(ValueError, match="no runs match"):
        plan_workbook(wb(sheet("A", ["zzz"])), ["a"])
    with pytest.raises(ValueError, match="both"):
        plan_workbook(wb(sheet("A", ["a"]), sheet("B", ["a*"])), ["a"])
    with pytest.raises(ValueError, match="not on any sheet"):
        plan_workbook(wb(sheet("A", ["a"]), unassigned="error"), ["a", "b"])
    assert [p.name for p in plan_workbook(wb(sheet("A", ["a"]), unassigned="ignore"), ["a", "b"])] == ["A"]


def test_spec_validation():
    with pytest.raises(ValidationError):
        sheet("A", ["a"], mode="aggregate")
    with pytest.raises(ValidationError):
        sheet("bad/name", ["a"])
    with pytest.raises(ValidationError):
        wb(sheet("A", ["a"]), sheet("a", ["b"]))


def test_merge_outer_join_with_underscore_suffix():
    a = pd.DataFrame({"m": ["Jan", "Feb"], "v": [1, 2]})
    b = pd.DataFrame({"m": ["Feb", "Mar"], "v": [3, 4]})
    out = merge_frames({"g/a": a, "g/b": b}, ["m"])
    assert list(out.columns) == ["m", "v_a", "v_b"]
    assert out["m"].tolist() == ["Feb", "Jan", "Mar"]


def test_workbook_round_trip(tmp_path):
    make_batch(tmp_path, {
        "g/a": ("m,v\nJan,1\nFeb,2\n", "ok"),
        "g/b": ("m,v\nFeb,3\nMar,4\n", "suspect"),
        "g/dup": ("m,v\nJan,1\nJan,2\n", "ok"),
        "ids": ("id,n\n007,1\n", "ok"),
        "bad": ("", "error"),
    })
    labels = ["g/a", "g/b", "g/dup", "ids", "bad"]
    spec = wb(sheet("G", ["g/*"], mode="aggregate", key=["m"]))
    (path,) = build_workbooks([spec], labels, tmp_path).paths
    book = load_workbook(path)
    assert book.sheetnames == ["G", "g_dup", "ids", "_notes"]
    assert [c.value for c in book["G"][1]] == ["m", "v_a", "v_b"]
    assert book["ids"]["A2"].value == "007"
    assert {r[0].value: r[1].value for r in book["_notes"].iter_rows(min_row=2)} == {
        "g/dup": "own sheet", "bad": "omitted", "g/b": "suspect",
    }


def test_blocks_layout_side_by_side(tmp_path):
    make_batch(tmp_path, {"a": ("m,v\nJan,1\n", "ok"), "b": ("m,v\nJan,2\nFeb,3\n", "ok")})
    spec = wb(sheet("S", ["*"], mode="aggregate", key=["m"], layout="blocks"))
    (path,) = build_workbooks([spec], ["a", "b"], tmp_path).paths
    ws = load_workbook(path)["S"]
    assert [c.value for c in ws[1]][:5] == ["a", None, None, "b", None]
    assert len(ws.tables) == 2


def test_missing_key_raises(tmp_path):
    make_batch(tmp_path, {"a": ("m,v\nJan,1\n", "ok"), "b": ("x,v\n1,2\n", "ok")})
    with pytest.raises(ValueError, match="no key"):
        build_workbooks([wb(sheet("S", ["*"], mode="aggregate", key=["m"]))], ["a", "b"], tmp_path)


def test_remove_csvs_keeps_manifest(tmp_path):
    make_batch(tmp_path, {"g/a": ("m\n1\n", "ok")})
    remove_csvs(tmp_path, ["g/a"])
    assert [p.name for p in tmp_path.iterdir()] == ["manifest.csv"]


def test_unassigned_ignore_omits_runs(tmp_path):
    make_batch(tmp_path, {"a": ("m\n1\n", "ok"), "b": ("m\n2\n", "ok")})
    result = build_workbooks([wb(sheet("A", ["a"]), unassigned="ignore")], ["a", "b"], tmp_path)
    assert load_workbook(result.paths[0]).sheetnames == ["A"]
    assert result.consumed == {"a"}


def test_unassigned_error_raises(tmp_path):
    make_batch(tmp_path, {"a": ("m\n1\n", "ok"), "b": ("m\n2\n", "ok")})
    with pytest.raises(ValueError, match="not on any sheet"):
        build_workbooks([wb(sheet("A", ["a"]), unassigned="error")], ["a", "b"], tmp_path)


def test_clean_keeps_csvs_not_in_workbook(tmp_path):
    make_batch(tmp_path, {"a": ("m\n1\n", "ok"), "b": ("m\n2\n", "ok"), "c": ("", "error")})
    result = build_workbooks([wb(sheet("A", ["a"]), unassigned="ignore")], ["a", "b", "c"], tmp_path)
    remove_csvs(tmp_path, result.consumed)
    assert not (tmp_path / "a.csv").exists()
    assert (tmp_path / "b.csv").exists()
    assert (tmp_path / "c.csv").exists()


PARTIAL_RUNS = {"r/a": ("m\n1\n", "ok"), "r/b": ("m\n2\n", "ok"), "s/c": ("m\n3\n", "ok"), "s/d": ("", "error")}
PARTIAL_LABELS = list(PARTIAL_RUNS)
PARTIAL_SPEC = wb(sheet("R", ["r/*"], mode="aggregate", key=["m"]), sheet("S", ["s/*"]), unassigned="ignore")


def test_partial_writes_partial_file_and_keeps_main(tmp_path):
    make_batch(tmp_path, PARTIAL_RUNS)
    (main,) = build_workbooks([PARTIAL_SPEC], PARTIAL_LABELS, tmp_path).paths
    before = main.read_bytes()
    (partial,) = build_workbooks([PARTIAL_SPEC], PARTIAL_LABELS, tmp_path, prefix="r/").paths
    assert partial.name == "w.partial.xlsx"
    assert main.read_bytes() == before


def test_partial_overwrites_prior_partial(tmp_path):
    make_batch(tmp_path, PARTIAL_RUNS)
    build_workbooks([PARTIAL_SPEC], PARTIAL_LABELS, tmp_path, prefix="r/")
    (partial,) = build_workbooks([PARTIAL_SPEC], PARTIAL_LABELS, tmp_path, prefix="s/").paths
    assert load_workbook(partial).sheetnames == ["s_c", "_notes"]


def test_partial_skips_out_of_scope_sheet_with_warning(tmp_path, caplog):
    make_batch(tmp_path, PARTIAL_RUNS)
    (partial,) = build_workbooks([PARTIAL_SPEC], PARTIAL_LABELS, tmp_path, prefix="r/").paths
    assert load_workbook(partial).sheetnames == ["R"]
    assert "sheet 'S' skipped" in caplog.text


def test_partial_scope_comes_from_manifest_and_notes_only_in_scope(tmp_path):
    make_batch(tmp_path, PARTIAL_RUNS)
    (partial,) = build_workbooks([PARTIAL_SPEC], PARTIAL_LABELS, tmp_path, prefix="s/").paths
    notes = load_workbook(partial)["_notes"]
    assert [r[0].value for r in notes.iter_rows(min_row=2)] == ["s/d"]


def test_partial_still_validates_full_config(tmp_path):
    make_batch(tmp_path, PARTIAL_RUNS)
    spec = wb(sheet("R", ["r/*"]), sheet("X", ["zzz"]))
    with pytest.raises(ValueError, match="no runs match"):
        build_workbooks([spec], PARTIAL_LABELS, tmp_path, prefix="r/")
